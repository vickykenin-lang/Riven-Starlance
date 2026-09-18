from __future__ import annotations

import io
import json
import os
import re
from datetime import datetime, timezone
from enum import StrEnum
from pathlib import Path
from uuid import UUID, uuid4

from docx import Document as DocxDocument
from pydantic import BaseModel, Field
from pypdf import PdfReader


MAX_UPLOAD_BYTES = 15 * 1024 * 1024
SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt"}
CHUNK_SIZE = 1800
CHUNK_OVERLAP = 250
DEFAULT_DATA_DIR = Path(os.getenv("RIVEN_DATA_DIR", "/tmp/riven-data"))


class DocumentStatus(StrEnum):
    PARSED = "parsed"
    FAILED = "failed"


class DocumentChunk(BaseModel):
    id: UUID = Field(default_factory=uuid4)
    document_id: UUID
    index: int
    text: str


class ResearchDocument(BaseModel):
    id: UUID = Field(default_factory=uuid4)
    filename: str
    content_type: str | None = None
    size_bytes: int
    uploaded_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))
    status: DocumentStatus
    char_count: int = 0
    chunk_count: int = 0
    error: str | None = None
    stored_path: str | None = None


class SourceContext(BaseModel):
    document_id: UUID
    filename: str
    chunk_id: UUID
    chunk_index: int
    text: str
    score: float


class DocumentStore:
    def __init__(self, data_dir: Path | str = DEFAULT_DATA_DIR) -> None:
        self.data_dir = Path(data_dir)
        self.files_dir = self.data_dir / "files"
        self.registry_path = self.data_dir / "registry.json"
        self.data_dir.mkdir(parents=True, exist_ok=True)
        self.files_dir.mkdir(parents=True, exist_ok=True)
        self._documents: dict[UUID, ResearchDocument] = {}
        self._chunks: dict[UUID, list[DocumentChunk]] = {}
        self._load()

    def _load(self) -> None:
        if not self.registry_path.exists():
            return
        try:
            payload = json.loads(self.registry_path.read_text(encoding="utf-8"))
            for raw in payload.get("documents", []):
                item = ResearchDocument.model_validate(raw)
                self._documents[item.id] = item
            for document_id, raw_chunks in payload.get("chunks", {}).items():
                self._chunks[UUID(document_id)] = [DocumentChunk.model_validate(chunk) for chunk in raw_chunks]
        except Exception:
            self._documents = {}
            self._chunks = {}

    def _persist(self) -> None:
        payload = {
            "documents": [item.model_dump(mode="json") for item in self._documents.values()],
            "chunks": {
                str(document_id): [chunk.model_dump(mode="json") for chunk in chunks]
                for document_id, chunks in self._chunks.items()
            },
        }
        temp_path = self.registry_path.with_suffix(".tmp")
        temp_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        temp_path.replace(self.registry_path)

    def list(self) -> list[ResearchDocument]:
        return sorted(self._documents.values(), key=lambda item: item.uploaded_at, reverse=True)

    def get(self, document_id: UUID) -> ResearchDocument | None:
        return self._documents.get(document_id)

    def chunks(self, document_id: UUID) -> list[DocumentChunk]:
        return list(self._chunks.get(document_id, []))

    def add(self, filename: str, content_type: str | None, payload: bytes) -> ResearchDocument:
        if len(payload) > MAX_UPLOAD_BYTES:
            raise ValueError("Document exceeds 15 MB upload limit")

        suffix = Path(filename).suffix.lower()
        if suffix not in SUPPORTED_SUFFIXES:
            raise ValueError("Unsupported document type. Use PDF, DOCX or TXT.")

        item = ResearchDocument(
            filename=filename,
            content_type=content_type,
            size_bytes=len(payload),
            status=DocumentStatus.PARSED,
        )
        safe_name = f"{item.id}{suffix}"
        stored_path = self.files_dir / safe_name
        stored_path.write_bytes(payload)
        item.stored_path = str(stored_path)

        try:
            text = extract_text(suffix, payload)
            if not text.strip():
                raise ValueError("No readable text found in document")
            chunks = chunk_text(item.id, text)
            item.char_count = len(text)
            item.chunk_count = len(chunks)
            self._chunks[item.id] = chunks
        except Exception as exc:
            item.status = DocumentStatus.FAILED
            item.error = str(exc)
            self._chunks[item.id] = []
        self._documents[item.id] = item
        self._persist()
        return item

    def retrieve(self, query: str, *, limit: int = 8) -> list[SourceContext]:
        query_terms = tokenize(query)
        ranked: list[tuple[float, ResearchDocument, DocumentChunk]] = []
        for document in self._documents.values():
            if document.status != DocumentStatus.PARSED:
                continue
            for chunk in self._chunks.get(document.id, []):
                terms = tokenize(chunk.text)
                if not terms:
                    continue
                overlap = len(query_terms & terms)
                score = overlap / max(1, len(query_terms))
                if overlap:
                    ranked.append((score, document, chunk))

        ranked.sort(key=lambda row: (-row[0], row[2].index))
        return [
            SourceContext(
                document_id=document.id,
                filename=document.filename,
                chunk_id=chunk.id,
                chunk_index=chunk.index,
                text=chunk.text,
                score=round(score, 4),
            )
            for score, document, chunk in ranked[:limit]
        ]


def tokenize(text: str) -> set[str]:
    return {token for token in re.findall(r"[A-Za-z0-9_]{3,}", text.lower())}


def extract_text(suffix: str, payload: bytes) -> str:
    if suffix == ".txt":
        return payload.decode("utf-8", errors="replace")
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(payload))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    if suffix == ".docx":
        document = DocxDocument(io.BytesIO(payload))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    raise ValueError(f"Unsupported suffix: {suffix}")


def chunk_text(document_id: UUID, text: str) -> list[DocumentChunk]:
    normalized = re.sub(r"\n{3,}", "\n\n", text).strip()
    chunks: list[DocumentChunk] = []
    start = 0
    index = 0
    while start < len(normalized):
        end = min(len(normalized), start + CHUNK_SIZE)
        if end < len(normalized):
            boundary = normalized.rfind("\n", start, end)
            if boundary > start + CHUNK_SIZE // 2:
                end = boundary
        fragment = normalized[start:end].strip()
        if fragment:
            chunks.append(DocumentChunk(document_id=document_id, index=index, text=fragment))
            index += 1
        if end >= len(normalized):
            break
        start = max(start + 1, end - CHUNK_OVERLAP)
    return chunks


document_store = DocumentStore()
